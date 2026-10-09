"""Bounded Office Open XML and PDF tools. No Office automation server needed."""
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import *
import io
import itertools
import subprocess
import tempfile


OOXML = {".docx", ".xlsx", ".pptx"}


def source(args):
    p = path(args["path"], suffixes=OOXML | {".pdf"})
    if p.suffix.lower() in OOXML:
        zip_guard(p)
    return p


def doc_paragraphs(doc):
    def table_paragraphs(table):
        seen = set()
        for row in table.rows:
            for cell in row.cells:
                if cell._tc in seen:
                    continue
                seen.add(cell._tc)
                yield from cell.paragraphs
                for nested in cell.tables:
                    yield from table_paragraphs(nested)
    yield from doc.paragraphs
    for table in doc.tables:
        yield from table_paragraphs(table)
    seen_parts = set()
    for section in doc.sections:
        for part in (section.header, section.footer, section.first_page_header,
                     section.first_page_footer, section.even_page_header, section.even_page_footer):
            # Linked sections were already visited. Accessing an undefined
            # header's paragraphs would create an unwanted empty header part.
            if part.is_linked_to_previous or part.part in seen_parts:
                continue
            seen_parts.add(part.part)
            yield from part.paragraphs
            for table in part.tables:
                yield from table_paragraphs(table)


def slide_shapes(shapes, depth=0):
    require(depth <= 32, "PowerPoint group nesting exceeds limit")
    for shape in shapes:
        yield shape
        if hasattr(shape, "shapes"):
            yield from slide_shapes(shape.shapes, depth+1)


def inspect(args):
    fields(args, ("path", "start", "limit", "sheet", "range", "data_only"), ("path",))
    p = source(args)
    start = integer(args.get("start", 0), "start", 0, 1000000)
    limit = integer(args.get("limit", 20), "limit", 1, 100)
    result = {"source": receipt(p), "format": p.suffix.lower()[1:]}
    if p.suffix.lower() == ".docx":
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        doc = Document(p)
        items = []
        blocks = list(doc.iter_inner_content())
        for index, block in enumerate(blocks[start:start + limit], start):
            if isinstance(block, Paragraph):
                items.append({"index": index, "kind": "paragraph", "text": block.text[:8000],
                              "style": block.style.name})
            elif isinstance(block, Table):
                items.append({"index": index, "kind": "table", "rows":
                              [[c.text[:500] for c in r.cells[:20]] for r in block.rows[:30]],
                              "rows_total": len(block.rows), "columns_total": len(block.columns)})
        result.update(items=items, total=len(blocks), next_start=start + limit if start + limit < len(blocks) else None)
    elif p.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        from openpyxl.utils.cell import range_boundaries
        require(type(args.get("data_only", False)) is bool, "data_only must be boolean")
        wb = load_workbook(p, read_only=True, data_only=args.get("data_only", False), keep_links=False)
        try:
            sheet = args.get("sheet", wb.sheetnames[0])
            require(sheet in wb.sheetnames, "unknown sheet")
            ws = wb[sheet]
            result.update(sheets=wb.sheetnames, sheet=sheet,
                          max_row=ws.max_row, max_column=ws.max_column,
                          formulas_evaluated=False)
            if "range" in args:
                string(args["range"], "range", 40)
                c1, r1, c2, r2 = range_boundaries(args["range"])
                require(all(type(n) is int and n > 0 for n in (c1, r1, c2, r2)) and
                        c2 <= 16384 and r2 <= 1048576 and c1 <= c2 and r1 <= r2 and
                        (c2-c1+1)*(r2-r1+1) <= 2000, "range must contain at most 2000 cells")
            else:
                c1, r1, c2, r2 = 1, start + 1, min(ws.max_column or 1, 20), start + limit
            result["rows"] = [[v if not isinstance(v, str) else v[:2000] for v in row]
                              for row in ws.iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2, values_only=True)]
        finally:
            wb.close()
    elif p.suffix.lower() == ".pptx":
        from pptx import Presentation
        prs = Presentation(p)
        slides = []
        for index, slide in enumerate(list(prs.slides)[start:start + limit], start):
            shapes = list(slide_shapes(slide.shapes))
            slides.append({"index": index, "text": [s.text[:8000] for s in shapes if s.has_text_frame],
                           "notes": slide.notes_slide.notes_text_frame.text[:8000] if slide.has_notes_slide else "",
                           "tables": [[[c.text[:500] for c in itertools.islice(row.cells,20)] for row in itertools.islice(s.table.rows,30)]
                                      for s in shapes if s.has_table],
                           "table_dimensions": [{"rows":len(s.table.rows),"columns":len(s.table.columns),
                                                 "truncated":len(s.table.rows)>30 or len(s.table.columns)>20}
                                                for s in shapes if s.has_table]})
        result.update(slides=slides, total=len(prs.slides),
                      next_start=start + limit if start + limit < len(prs.slides) else None)
    else:
        from pypdf import PdfReader
        reader = PdfReader(p)
        require(not reader.is_encrypted, "encrypted PDF requires operator decryption first")
        pages = []
        for index in range(start, min(start + limit, len(reader.pages))):
            page = reader.pages[index]
            stream = page.get_contents()
            require(stream is None or len(stream.get_data()) <= 16 * 1024 * 1024, "PDF page content too large")
            text = page.extract_text(extraction_mode="layout") or ""
            pages.append({"index": index, "text": text[:16000], "truncated": len(text) > 16000,
                          "ocr_suggested": len(text.strip()) < 30})
        result.update(pages=pages, total=len(reader.pages),
                      fields={k: {"type": str(v.get("/FT", "")), "value": str(v.get("/V", ""))[:1000]}
                              for k, v in (reader.get_fields() or {}).items()},
                      next_start=start + limit if start + limit < len(reader.pages) else None)
    return result


def spec(args):
    p = path(args["spec"], suffixes={".json"}, max_bytes=2 * 1024 * 1024)
    value = json.loads(p.read_text())
    require(isinstance(value, dict), "spec must be an object")
    return value


def blocks(doc, items):
    require(isinstance(items, list) and len(items) <= 1000, "blocks must be a bounded list")
    for item in items:
        fields(item, ("kind", "text", "level", "rows", "path", "width", "runs"), ("kind",))
        kind = item["kind"]
        if kind in ("paragraph", "heading", "bullet", "number"):
            text = string(item.get("text", ""), "block text", 10000, empty=True)
            if kind == "heading":
                p = doc.add_heading(text, integer(item.get("level", 1), "heading level", 0, 9))
            else:
                style = {"paragraph": None, "bullet": "List Bullet", "number": "List Number"}[kind]
                p = doc.add_paragraph(text, style=style)
            if "runs" in item:
                require(not text and isinstance(item["runs"], list) and len(item["runs"]) <= 100, "use text or runs")
                for value in item["runs"]:
                    fields(value, ("text", "bold", "italic"), ("text",))
                    run = p.add_run(string(value["text"], "run text", 10000, empty=True))
                    for key in ("bold", "italic"):
                        if key in value:
                            require(type(value[key]) is bool, "run styling must be boolean")
                            setattr(run, key, value[key])
        elif kind == "table":
            rows = item.get("rows")
            require(isinstance(rows, list) and 0 < len(rows) <= 500 and
                    all(isinstance(r, list) and len(r) == len(rows[0]) for r in rows) and
                    0 < len(rows[0]) <= 30 and len(rows)*len(rows[0]) <= 5000, "invalid table")
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for r, values in zip(table.rows, rows):
                for cell, value in zip(r.cells, values):
                    require(value is None or type(value) in (str, int, float, bool), "table cells must be text or scalar values")
                    if type(value) is float:
                        number(value, "table number", -1e300, 1e300)
                    cell.text = string("" if value is None else str(value), "table text", 10000, empty=True)
        elif kind == "image":
            from docx.shared import Inches
            doc.add_picture(str(path(item["path"], suffixes={".png", ".jpg", ".jpeg"})),
                            width=Inches(number(item.get("width", 5), "image width", .1, 7)))
        elif kind == "page_break":
            doc.add_page_break()
        else:
            raise ToolError("unsupported document block")


def compose(args):
    fields(args, ("format", "output", "spec"), ("format", "output", "spec"))
    kind = args["format"]
    require(kind in ("docx", "xlsx", "pptx", "pdf"), "unsupported output format")
    data = spec(args)
    with output(args["output"], {"." + kind}) as temp:
        if kind == "docx":
            from docx import Document
            doc = Document()
            fields(data, ("title", "blocks"), ("blocks",))
            if data.get("title"):
                doc.core_properties.title = string(data["title"], "title", 255)
            blocks(doc, data["blocks"])
            doc.save(temp)
        elif kind == "xlsx":
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
            from openpyxl.worksheet.table import Table, TableStyleInfo
            fields(data, ("sheets",), ("sheets",))
            require(isinstance(data["sheets"], list) and 0 < len(data["sheets"]) <= 30, "invalid sheets")
            wb = Workbook()
            wb.remove(wb.active)
            titles = set()
            for sheet in data["sheets"]:
                fields(sheet, ("name", "rows", "freeze", "widths", "table"), ("name", "rows"))
                name = string(sheet["name"], "sheet name", 31)
                require(not re.search(r'[\\/*?:\[\]]', name) and name.lower() not in titles, "invalid or duplicate sheet name")
                titles.add(name.lower())
                ws = wb.create_sheet(name)
                require(isinstance(sheet["rows"], list) and len(sheet["rows"]) <= 10000, "too many rows")
                require(sum(len(r) for r in sheet["rows"] if isinstance(r, list)) <= 100000, "too many cells")
                for row in sheet["rows"]:
                    require(isinstance(row, list) and len(row) <= 100, "invalid row")
                    ws.append([cell_value(v) for v in row])
                    # Plain strings beginning '=' are literal; formula objects are explicit.
                    for col, value in enumerate(row, 1):
                        if isinstance(value, str):
                            ws.cell(ws.max_row, col).data_type = "s"
                if sheet.get("freeze"):
                    ws.freeze_panes = cell_ref(sheet["freeze"])
                for col, width in sheet.get("widths", {}).items():
                    require(re.fullmatch(r"[A-Z]{1,3}", col), "invalid width column")
                    ws.column_dimensions[col].width = number(width, "column width", 1, 100)
                if sheet.get("table"):
                    require(ws.max_row >= 2 and all(isinstance(c.value, str) and c.value for c in ws[1]), "table needs text headers and data")
                    t = Table(displayName=f"Table{len(titles)}", ref=ws.dimensions)
                    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
                    ws.add_table(t)
            wb.save(temp)
            wb.close()
        elif kind == "pptx":
            from pptx import Presentation
            from pptx.util import Inches, Pt
            fields(data, ("slides", "widescreen"), ("slides",))
            require(isinstance(data["slides"], list) and 0 < len(data["slides"]) <= 100, "invalid slides")
            prs = Presentation()
            if data.get("widescreen", True):
                prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
            for slide_data in data["slides"]:
                fields(slide_data, ("title", "bullets", "notes", "image"), ("title",))
                slide = prs.slides.add_slide(prs.slide_layouts[1])
                slide.shapes.title.text = string(slide_data["title"], "slide title", 300)
                bullets = slide_data.get("bullets", [])
                require(isinstance(bullets, list) and len(bullets) <= 12, "too many slide bullets")
                tf = slide.placeholders[1].text_frame
                tf.clear()
                for i, text in enumerate(bullets):
                    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    p.text = string(text, "bullet", 1000)
                    p.font.size = Pt(24)
                if "image" in slide_data:
                    slide.shapes.add_picture(str(path(slide_data["image"], suffixes={".png", ".jpg", ".jpeg"})),
                                             Inches(8), Inches(2), width=Inches(4.5))
                if "notes" in slide_data:
                    slide.notes_slide.notes_text_frame.text = string(slide_data["notes"], "notes", 10000, empty=True)
            prs.save(temp)
        else:
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.ttfonts import TTFont
            from xml.sax.saxutils import escape
            fields(data, ("title", "blocks", "font"), ("blocks", "font"))
            # A user-supplied font is required, so Vietnamese never falls back to missing glyphs.
            font = path(data["font"], suffixes={".ttf"})
            pdfmetrics.registerFont(TTFont("SootDocument", str(font)))
            styles = getSampleStyleSheet()
            for s in styles.byName.values():
                s.fontName = "SootDocument"
            story = []
            require(isinstance(data["blocks"], list) and len(data["blocks"]) <= 1000, "invalid blocks")
            for b in data["blocks"]:
                fields(b, ("kind", "text", "rows"), ("kind",))
                if b["kind"] in ("heading", "paragraph", "bullet"):
                    text = escape(string(b.get("text", ""), "PDF text", 10000, empty=True)).replace("\n", "<br/>")
                    story.append(Paragraph(text, styles["Heading1" if b["kind"] == "heading" else "BodyText"]))
                    story.append(Spacer(1, 8))
                elif b["kind"] == "table":
                    rows = b.get("rows")
                    require(isinstance(rows, list) and 0 < len(rows) <= 500 and
                            all(isinstance(r, list) and 0 < len(r) <= 20 for r in rows), "invalid PDF table")
                    table = Table([[Paragraph(escape(str(v)), styles["BodyText"]) for v in r] for r in rows], repeatRows=1)
                    table.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))
                    story.append(table)
                elif b["kind"] == "page_break":
                    story.append(PageBreak())
                else:
                    raise ToolError("unsupported PDF block")
            require(story, "PDF needs content")
            SimpleDocTemplate(str(temp), title=data.get("title", ""), leftMargin=48, rightMargin=48).build(story)
    return receipt(path(args["output"]), format=kind, formulas_evaluated=False if kind == "xlsx" else None)


def cell_ref(value):
    require(isinstance(value, str) and re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]{0,6}", value), "invalid cell reference")
    from openpyxl.utils.cell import coordinate_to_tuple
    r, c = coordinate_to_tuple(value)
    require(r <= 1048576 and c <= 16384, "cell outside Excel limits")
    return value


def cell_value(value):
    require(value is None or type(value) in (str, int, float, bool, dict), "invalid cell value")
    if isinstance(value, dict):
        fields(value, ("formula",), ("formula",))
        f = string(value["formula"], "formula", 8192)
        require(f.startswith("=") and not re.search(r"(?i)(WEBSERVICE|HYPERLINK|DDE|RTD|\[|\||https?://)", f),
                "formula contains an external operation or workbook link")
        return f
    if isinstance(value, str):
        string(value, "cell text", 32767, empty=True)
    if type(value) is float:
        number(value, "cell number", -1e300, 1e300)
    return value


def replace_runs(paragraph, old, new):
    # Replace spans even when Word/PPT split a template token across styled runs.
    runs = paragraph.runs
    text = "".join(r.text for r in runs)
    matches = list(re.finditer(re.escape(old), text))
    for match in reversed(matches):
        offset = 0
        first = True
        for run in runs:
            end = offset + len(run.text)
            a, b = max(match.start(), offset), min(match.end(), end)
            if a < b:
                run.text = run.text[:a-offset] + (new if first else "") + run.text[b-offset:]
                first = False
            offset = end
    return len(matches)


def edit(args):
    fields(args, ("path", "output", "spec"), ("path", "output", "spec"))
    p = source(args)
    require(p.suffix.lower() in OOXML, "use pdf_modify for PDF changes")
    data = spec(args)
    with output(args["output"], {p.suffix.lower()}) as temp:
        if p.suffix.lower() == ".xlsx":
            from openpyxl import load_workbook
            from openpyxl.styles import Font, Alignment
            fields(data, ("cells",), ("cells",))
            cells = data["cells"]
            require(isinstance(cells, list) and 0 < len(cells) <= 5000, "invalid cell changes")
            wb = load_workbook(p, keep_links=False)
            for change in cells:
                fields(change, ("sheet", "cell", "value", "number_format", "bold"), ("sheet", "cell", "value"))
                require(change["sheet"] in wb.sheetnames, "unknown sheet")
                c = wb[change["sheet"]][cell_ref(change["cell"])]
                c.value = cell_value(change["value"])
                if isinstance(change["value"], str):
                    c.data_type = "s"
                if "number_format" in change:
                    c.number_format = string(change["number_format"], "number format", 100)
                if "bold" in change:
                    import copy
                    require(type(change["bold"]) is bool, "bold must be boolean")
                    f = copy.copy(c.font)
                    f.bold = change["bold"]
                    c.font = f
            wb.save(temp)
            wb.close()
            count = len(cells)
        else:
            fields(data, ("replacements",), ("replacements",))
            values = data["replacements"]
            require(isinstance(values, list) and 0 < len(values) <= 100, "invalid replacements")
            if p.suffix.lower() == ".docx":
                from docx import Document
                doc = Document(p)
                paragraphs = list(doc_paragraphs(doc))
            else:
                from pptx import Presentation
                doc = Presentation(p)
                paragraphs = []
                for slide in doc.slides:
                    for shape in slide_shapes(slide.shapes):
                        if shape.has_text_frame:
                            paragraphs.extend(shape.text_frame.paragraphs)
                        if shape.has_table:
                            paragraphs.extend(q for row in shape.table.rows for cell in row.cells for q in cell.text_frame.paragraphs)
                    if slide.has_notes_slide:
                        paragraphs.extend(slide.notes_slide.notes_text_frame.paragraphs)
            count = 0
            for value in values:
                fields(value, ("find", "replace", "expected"), ("find", "replace"))
                old = string(value["find"], "find", 1000)
                new = string(value["replace"], "replacement", 10000, empty=True)
                n = sum(replace_runs(q, old, new) for q in paragraphs)
                if "expected" in value:
                    require(n == integer(value["expected"], "expected matches", 0, 10000), "replacement count differs from expected; no output published")
                count += n
            doc.save(temp)
    return receipt(path(args["output"]), changes=count)


def pdf_modify(args):
    fields(args, ("paths", "output", "pages", "rotation", "fields"), ("paths", "output"))
    from pypdf import PdfReader, PdfWriter
    require(isinstance(args["paths"], list) and 0 < len(args["paths"]) <= 20, "paths needs 1..20 PDFs")
    writer = PdfWriter()
    for name in args["paths"]:
        reader = PdfReader(path(name, suffixes={".pdf"}))
        require(not reader.is_encrypted and len(reader.pages) <= 5000, "encrypted or oversized PDF")
        pages = args.get("pages")
        if pages is not None:
            require(len(args["paths"]) == 1 and isinstance(pages, list) and 0 < len(pages) <= 1000,
                    "page selection requires one source; pages are zero-based")
            for i in pages:
                integer(i, "page index", 0, len(reader.pages)-1)
        writer.append(reader, pages=pages)
    if "rotation" in args:
        require(type(args["rotation"]) is int and args["rotation"] in (0, 90, 180, 270), "invalid rotation")
        for page in writer.pages:
            page.rotate(args["rotation"])
    if "fields" in args:
        values = args["fields"]
        known = writer.get_fields() or {}
        require(isinstance(values, dict) and 0 < len(values) <= 100 and set(values).issubset(known), "unknown form fields")
        for value in values.values():
            string(value, "field value", 2000, empty=True)
        for page in writer.pages:
            if page.get("/Annots"):
                writer.update_page_form_field_values(page, values, auto_regenerate=False)
    require(0 < len(writer.pages) <= 5000, "merged PDF too large or empty")
    with output(args["output"], {".pdf"}) as temp:
        writer.write(temp)
    return receipt(path(args["output"]), pages=len(writer.pages), flattened=False)


def render(args):
    fields(args, ("path", "output", "page", "scale"), ("path", "output"))
    import pypdfium2 as pdfium
    p = path(args["path"], suffixes={".pdf"})
    scale = number(args.get("scale", 1.5), "scale", .25, 3)
    with pdfium.PdfDocument(p) as doc:
        page_number = integer(args.get("page", 0), "page", 0, len(doc)-1)
        with contextlib.closing(doc[page_number]) as page:
            w, h = page.get_size()
            require(w*h*scale*scale <= 32_000_000, "render exceeds pixel limit")
            bitmap = page.render(scale=scale)
            try:
                with output(args["output"], {".png"}) as temp:
                    bitmap.to_pil().save(temp)
            finally:
                bitmap.close()
    return receipt(path(args["output"]), page=page_number)


def convert(args):
    fields(args, ("path", "output"), ("path", "output"))
    p = path(args["path"], suffixes=OOXML | {".doc", ".xls", ".ppt", ".odt", ".ods", ".odp"})
    destination = path(args["output"], existing=False, suffixes=OOXML | {".pdf"})
    config = binding("OFFICE_BINDING")
    binary = Path(config.get("libreoffice", ""))
    require(binary.is_absolute() and binary.is_file(), "host must bind an absolute LibreOffice executable")
    with tempfile.TemporaryDirectory(prefix="soot-office-") as directory:
        sandbox = Path(directory)
        profile = (sandbox / "profile").as_uri()
        result = subprocess.run([str(binary), "-env:UserInstallation=" + profile, "--headless", "--convert-to",
                                 destination.suffix.lower()[1:], "--outdir", directory, str(p)],
                                capture_output=True, timeout=120, env={"HOME": directory, "PATH": "/usr/bin:/bin"})
        converted = sandbox / (p.stem + destination.suffix.lower())
        require(result.returncode == 0 and converted.is_file(), "LibreOffice conversion failed")
        with output(args["output"], {destination.suffix.lower()}) as temp:
            temp.write_bytes(converted.read_bytes())
    return receipt(path(args["output"]), renderer="libreoffice", visual_qa="pending")


if __name__ == "__main__":
    invoke("office", {"inspect": inspect, "compose": compose, "edit": edit,
                     "pdf_modify": pdf_modify, "render": render, "convert": convert})
