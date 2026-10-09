#!/usr/bin/env python3
"""Opt-in real file + LibreOffice journeys through the public pack wrappers."""
import argparse
import hashlib
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import time

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument("--output", type=Path, required=True, help="New evidence directory")
parser.add_argument("--libreoffice", type=Path, required=True)
parser.add_argument("--ocr-binding-file", type=Path, help="Explicit real DeepSeek OCR service credential JSON")
args = parser.parse_args()
root = args.output.resolve()
root.mkdir()  # Never overwrite previous evidence.
repo = Path(__file__).resolve().parent.parent
env = {"FILE_ROOT": str(root), "PYTHON_EXECUTABLE": sys.executable,
       "OFFICE_BINDING": json.dumps({"libreoffice": str(args.libreoffice.resolve())})}
report = {"status":"FAIL", "interpreter":sys.executable,"cases":[], "live_scan_ocr":"NOT RUN: no real DeepSeek OCR binding supplied"}


def call(pack, operation, input, fail=None):
    started=time.monotonic()
    envelope={"api_version":"soot/v1","hook":"tool.call","soot":"office-ocr-qa",
              "operation":f"{pack}_{operation}","input":input}
    result=subprocess.run([str(repo/"packs"/pack/"tools/run"),operation],input=json.dumps(envelope),
                          capture_output=True,text=True,env=env,timeout=200)
    if fail:
        assert result.returncode == 1, (operation,result.stderr)
        assert fail in result.stderr, result.stderr
        response=json.loads(result.stderr)
    else:
        assert result.returncode == 0, (operation,result.stderr)
        response=json.loads(result.stdout)
    report["cases"].append({"pack":pack,"operation":operation,"input":input,"exit_code":result.returncode,
                            "seconds":round(time.monotonic()-started,3),"result":response})
    return response


def spec(name, value):
    (root/name).write_text(json.dumps(value,ensure_ascii=False))
    return name


try:
    from docx import Document
    from openpyxl import load_workbook
    from pptx import Presentation
    from pypdf import PdfReader
    from PIL import Image
    from reportlab.pdfgen import canvas
    font=next((p for p in (Path('/System/Library/Fonts/Supplemental/Arial.ttf'),
                          Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')) if p.is_file()),None)
    assert font, "Local TrueType font required"
    shutil.copyfile(font,root/'font.ttf')
    content='Báo cáo tiếng Việt: Tổng ngân sách 1.650.000 đồng.'
    word=spec('word.json',{'blocks':[{'kind':'heading','text':'Báo cáo tiếng Việt','level':1},
                                   {'kind':'paragraph','text':content},
                                   {'kind':'table','rows':[['Tên','Số tiền'],['Máy chủ',1200000],['Phần mềm',450000]]}]})
    call('office','compose',{'format':'docx','spec':word,'output':'source.docx'})
    assert Document(root/'source.docx').tables[0].cell(1,0).text == 'Máy chủ'
    edit=spec('edit.json',{'replacements':[{'find':'Báo cáo','replace':'Kết quả','expected':2}]})
    original=hashlib.sha256((root/'source.docx').read_bytes()).hexdigest()
    call('office','edit',{'path':'source.docx','spec':edit,'output':'edited.docx'})
    assert Document(root/'edited.docx').paragraphs[0].text == 'Kết quả tiếng Việt'
    assert hashlib.sha256((root/'source.docx').read_bytes()).hexdigest() == original
    sheet=spec('sheet.json',{'sheets':[{'name':'Chi phí','rows':[['Mục','VND'],['Máy chủ',1200000],
                     ['Phần mềm',450000],['Tổng',{'formula':'=SUM(B2:B3)'}]]}]})
    call('office','compose',{'format':'xlsx','spec':sheet,'output':'source.xlsx'})
    book=load_workbook(root/'source.xlsx'); assert book['Chi phí']['B4'].value == '=SUM(B2:B3)'; book.close()
    slides=spec('slides.json',{'slides':[{'title':'Báo cáo tiếng Việt','bullets':[content], 'notes':'Tài liệu thử nghiệm.'}]})
    call('office','compose',{'format':'pptx','spec':slides,'output':'source.pptx'})
    assert Presentation(root/'source.pptx').slides[0].notes_slide.notes_text_frame.text == 'Tài liệu thử nghiệm.'
    pdf=spec('pdf.json',{'font':'font.ttf','blocks':[{'kind':'paragraph','text':content}]})
    call('office','compose',{'format':'pdf','spec':pdf,'output':'native.pdf'})
    assert content in PdfReader(root/'native.pdf').pages[0].extract_text()
    for format in ('docx','xlsx','pptx'):
        call('office','inspect',{'path':f'source.{format}'})
        call('office','convert',{'path':f'source.{format}','output':f'{format}.pdf'})
        assert len(PdfReader(root/f'{format}.pdf').pages) >= 1
        text=''.join(p.extract_text() for p in PdfReader(root/f'{format}.pdf').pages)
        assert ('1.650.000' if format != 'xlsx' else '1650000') in text.replace(',',''), text
        call('office','render',{'path':f'{format}.pdf','output':f'{format}.png'})
        with Image.open(root/f'{format}.png') as image: image.load(); assert len(image.convert('L').getcolors(image.width*image.height)) > 2
    call('office','pdf_modify',{'paths':['native.pdf','docx.pdf'],'output':'merged.pdf','rotation':90})
    assert len(PdfReader(root/'merged.pdf').pages) == 2
    call('ocr','assess',{'path':'native.pdf'})
    result=call('ocr','recognize',{'path':'native.pdf','output':'native.md'})
    assert result['pages'][0]['engine'] == 'native-text'
    assert content in (root/'native.md').read_text()
    call('office','render',{'path':'native.pdf','output':'scan.png'})
    scan=call('ocr','assess',{'path':'scan.png'}); assert scan['route'] == 'deepseek-ocr'
    call('ocr','recognize',{'path':'scan.png','output':'blocked-scan.md'},fail='host must bind OCR_BINDING')
    assert not (root/'blocked-scan.md').exists()
    call('office','compose',{'format':'docx','spec':word,'output':'source.docx'},fail='output must be new')
    call('office','inspect',{'path':'../escape.docx'},fail='path must be clean')
    if args.ocr_binding_file:
        env['OCR_BINDING']=args.ocr_binding_file.read_text()
        call('ocr','recognize',{'path':'scan.png','output':'scan.md'})
        actual=(root/'scan.md').read_text()
        assert '1.650.000' in actual and 'ngân sách' in actual, 'Live OCR failed ground-truth checks'
        report['live_scan_ocr']='PASS: real bound model recognized Vietnamese phrase and exact amount; remaining layout requires visual review'
    report['status']='PASS'
finally:
    (root/'report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
print(json.dumps({'status':report['status'],'cases':len(report['cases']),'live_scan_ocr':report['live_scan_ocr']}))
