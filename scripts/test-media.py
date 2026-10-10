#!/usr/bin/env python3
"""Real file/FFmpeg tests plus transport fixtures; no model or Google claims.

Run with the prepared Office pack interpreter. Missing dependencies, fonts or
FFmpeg are prerequisites that fail, never silently skipped. Network fixtures
replace only service transport; production validation/receipts still execute.
"""
import base64
from contextlib import contextmanager
import hashlib
import importlib.util
import io
import json
import math
import os
from pathlib import Path
import shutil
import struct
import subprocess
import sys
import tempfile
import threading
import time
import unittest
from unittest.mock import patch
import wave
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

REPO = Path(__file__).resolve().parent.parent
MODULES = {}
for pack in ("office", "ocr", "tts", "google-workspace", "video-editor"):
    spec = importlib.util.spec_from_file_location(pack.replace("-", "_"), REPO / "packs" / pack / "tools/runner.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    MODULES[pack] = module
office, ocr, tts, google, video = (MODULES[p] for p in ("office", "ocr", "tts", "google-workspace", "video-editor"))
from common import ToolError, Uncertain


def font_file():
    candidates = [Path("/System/Library/Fonts/Supplemental/Arial.ttf"),
                  Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")]
    found = next((p for p in candidates if p.is_file()), None)
    if found is None:
        raise RuntimeError("Test needs a local TrueType font (Arial on macOS or DejaVu Sans on Linux)")
    return found


def wav_bytes(duration=.25):
    data = io.BytesIO()
    with wave.open(data, "wb") as w:
        w.setnchannels(1)
        w.setsampwidth(2)
        w.setframerate(24000)
        w.writeframes(b"".join(struct.pack("<h", round(4000*math.sin(2*math.pi*440*i/24000)))
                              for i in range(round(duration*24000))))
    return data.getvalue()


@contextmanager
def service(responder):
    requests = []
    class Handler(BaseHTTPRequestHandler):
        def log_message(self, *args):
            pass
        def perform(self):
            raw = self.rfile.read(int(self.headers.get("Content-Length", "0")))
            requests.append((self.command, self.path, raw, dict(self.headers)))
            status, payload, headers = responder(requests[-1])
            self.send_response(status)
            for k, v in headers.items():
                self.send_header(k, v)
            self.send_header("Content-Length", str(len(payload)))
            self.end_headers()
            self.wfile.write(payload)
        do_GET = perform
        do_POST = perform
    server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        yield f"http://127.0.0.1:{server.server_port}", requests
    finally:
        server.shutdown()
        server.server_close()
        thread.join()


class MediaTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="forge-media-test-")
        self.root = Path(self.temp.name)
        self.env = patch.dict(os.environ, {"FILE_ROOT": str(self.root)})
        self.env.start()

    def tearDown(self):
        self.env.stop()
        self.temp.cleanup()

    def spec(self, value, name="spec.json"):
        (self.root / name).write_text(json.dumps(value, ensure_ascii=False))
        return name

    def font(self):
        shutil.copyfile(font_file(), self.root / "font.ttf")
        return "font.ttf"

    def pdf(self, name="source.pdf", text="Báo cáo tiếng Việt: một trăm hai mươi lăm triệu đồng."):
        spec = self.spec({"font": self.font(), "blocks": [{"kind": "paragraph", "text": text}]})
        office.compose({"format": "pdf", "spec": spec, "output": name})
        return name

    def test_path_boundary_traversal_absolute_hidden_and_symlink(self):
        outside = self.root.parent / "unrelated.txt"
        (self.root / "link").symlink_to(self.root.parent, target_is_directory=True)
        for value in ("../unrelated.txt", str(outside), ".private", "folder/../file", "link/unrelated.txt", "a//b"):
            with self.subTest(value=value), self.assertRaises(ToolError):
                office.path(value)

    def test_atomic_output_never_overwrites_and_cleans_failed_staging(self):
        (self.root / "existing.txt").write_text("original")
        with self.assertRaises(ToolError), office.output("existing.txt"):
            pass
        self.assertEqual((self.root / "existing.txt").read_text(), "original")
        with self.assertRaises(RuntimeError):
            with office.output("new.txt") as p:
                p.write_text("partial")
                raise RuntimeError("processing failed")
        self.assertFalse((self.root / "new.txt").exists())
        self.assertEqual(list(self.root.glob(".soot-*")), [])

    def test_public_envelope_wrapper_supports_forward_fields_and_rejects_mismatch(self):
        doc = self.spec({"blocks": [{"kind": "paragraph", "text": "Đúng hợp đồng"}]})
        call = {"api_version": "soot/v1", "hook": "tool.call", "soot": "test",
                "operation": "office_compose", "input": {"format": "docx", "spec": doc, "output": "out.docx"},
                "future_v1_field": {"allowed": True}}
        run = REPO / "packs/office/tools/run"
        result = subprocess.run([str(run), "compose"], input=json.dumps(call), text=True, capture_output=True,
                                env={"FILE_ROOT": str(self.root), "PYTHON_EXECUTABLE": sys.executable})
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(json.loads(result.stdout)["path"], "out.docx")
        call["operation"] = "office_edit"
        result = subprocess.run([str(run), "compose"], input=json.dumps(call), text=True, capture_output=True,
                                env={"FILE_ROOT": str(self.root), "PYTHON_EXECUTABLE": sys.executable})
        self.assertEqual(result.returncode, 1)
        self.assertIn("operation mismatch", result.stderr)

    def test_word_round_trip_document_order_table_and_vietnamese(self):
        from docx import Document
        s = self.spec({"blocks": [{"kind": "heading", "text": "Tiếng Việt", "level": 1},
                                 {"kind": "table", "rows": [["Tên", "Số"], ["Ngọc Anh", 125]]},
                                 {"kind": "paragraph", "text": "Kết thúc"}]})
        result = office.compose({"format": "docx", "spec": s, "output": "report.docx"})
        self.assertEqual(result["sha256"], hashlib.sha256((self.root/"report.docx").read_bytes()).hexdigest())
        actual = office.inspect({"path": "report.docx", "limit": 2})
        self.assertEqual([i["kind"] for i in actual["items"]], ["paragraph", "table"])
        self.assertEqual(actual["items"][1]["rows"][1], ["Ngọc Anh", "125"])
        self.assertEqual(actual["next_start"], 2)
        self.assertEqual(office.inspect({"path": "report.docx", "start": 2})["items"][0]["text"], "Kết thúc")

    def test_word_split_run_replacement_preserves_surrounding_style_and_headers(self):
        from docx import Document
        d = Document()
        p = d.add_paragraph()
        p.add_run("Trước ").italic = True
        p.add_run("{{na").bold = True
        p.add_run("me}} sau").italic = True
        d.sections[0].header.paragraphs[0].text = "{{name}}"
        d.add_table(rows=1, cols=1).cell(0,0).text = "{{name}}"
        d.save(self.root / "template.docx")
        s = self.spec({"replacements": [{"find": "{{name}}", "replace": "Ngọc Anh", "expected": 3}]})
        result = office.edit({"path": "template.docx", "spec": s, "output": "edited.docx"})
        self.assertEqual(result["changes"], 3)
        edited = Document(self.root / "edited.docx")
        self.assertEqual(edited.paragraphs[0].text, "Trước Ngọc Anh sau")
        self.assertTrue(edited.paragraphs[0].runs[0].italic)
        self.assertTrue(edited.paragraphs[0].runs[1].bold)
        self.assertEqual(edited.sections[0].header.paragraphs[0].text, "Ngọc Anh")
        self.assertEqual(Document(self.root / "template.docx").paragraphs[0].text, "Trước {{name}} sau")

    def test_replacement_count_guard_publishes_no_output(self):
        from docx import Document
        d = Document(); d.add_paragraph("source"); d.save(self.root / "source.docx")
        s = self.spec({"replacements": [{"find": "missing", "replace": "x", "expected": 1}]})
        with self.assertRaises(ToolError):
            office.edit({"path": "source.docx", "output": "fail.docx", "spec": s})
        self.assertFalse((self.root / "fail.docx").exists())

    def test_word_linked_first_and_even_headers_edit_once_without_creating_headers(self):
        from docx import Document
        d = Document()
        d.add_paragraph("Body {{name}}")
        d.sections[0].header.paragraphs[0].text = "Header {{name}}"
        d.sections[0].first_page_header.paragraphs[0].text = "First {{name}}"
        d.sections[0].even_page_header.paragraphs[0].text = "Even {{name}}"
        d.add_section()
        d.save(self.root / "headers.docx")
        s = self.spec({"replacements":[{"find":"{{name}}", "replace":"Ngọc Anh", "expected":4}]})
        office.edit({"path":"headers.docx", "spec":s, "output":"edited.docx"})
        edited = Document(self.root/'edited.docx')
        self.assertEqual(edited.sections[1].header.paragraphs[0].text, "Header Ngọc Anh")
        self.assertTrue(edited.sections[0].footer.is_linked_to_previous)
        self.assertEqual(edited.sections[0].first_page_header.paragraphs[0].text, "First Ngọc Anh")
        self.assertEqual(edited.sections[0].even_page_header.paragraphs[0].text, "Even Ngọc Anh")

    def test_uppercase_extensions_keep_format_routing_and_copy_edit(self):
        s = self.spec({"blocks":[{"kind":"paragraph", "text":"{{name}}"}]})
        office.compose({"format":"docx", "spec":s, "output":"SOURCE.DOCX"})
        self.assertEqual(office.inspect({"path":"SOURCE.DOCX"})["format"], "docx")
        s = self.spec({"replacements":[{"find":"{{name}}", "replace":"Ngọc Anh", "expected":1}]})
        office.edit({"path":"SOURCE.DOCX", "spec":s, "output":"EDITED.DOCX"})
        self.assertEqual(office.inspect({"path":"EDITED.DOCX"})["items"][0]["text"], "Ngọc Anh")
        self.pdf(name="SOURCE.PDF")
        self.assertEqual(ocr.assess({"path":"SOURCE.PDF"})["pages"][0]["route"], "native-text")

    def test_excel_formulas_are_explicit_literal_injection_and_cached_values(self):
        s = self.spec({"sheets": [{"name": "Chi phí", "rows": [["Mục", "VND"], ["A", 125],
                         ["B", {"formula": "=SUM(B2:B2)"}], ["=HYPERLINK(\"x\")", "literal"]], "freeze": "A2"}]})
        office.compose({"format": "xlsx", "spec": s, "output": "report.xlsx"})
        from openpyxl import load_workbook
        w = load_workbook(self.root / "report.xlsx")
        self.assertEqual(w.active["B3"].data_type, "f")
        self.assertEqual(w.active["A4"].data_type, "s")
        self.assertEqual(w.active.freeze_panes, "A2")
        w.close()
        actual = office.inspect({"path": "report.xlsx", "range": "A2:B3", "data_only": True})
        self.assertEqual(actual["rows"][1][1], None)
        self.assertFalse(actual["formulas_evaluated"])

    def test_excel_cell_copy_and_range_limits(self):
        s = self.spec({"sheets": [{"name": "Data", "rows": [["Count"], [4]]}]})
        office.compose({"format": "xlsx", "spec": s, "output": "a.xlsx"})
        s = self.spec({"cells": [{"sheet": "Data", "cell": "A2", "value": 7, "bold": True}]})
        office.edit({"path": "a.xlsx", "spec": s, "output": "b.xlsx"})
        self.assertEqual(office.inspect({"path": "b.xlsx", "range": "A2:A2"})["rows"], [[7]])
        self.assertEqual(office.inspect({"path": "a.xlsx", "range": "A2:A2"})["rows"], [[4]])
        for rng in ("A1:XFD1000", "A:A", "A0:B1", "XFE1:XFE2"):
            with self.subTest(rng=rng), self.assertRaises(Exception):
                office.inspect({"path": "a.xlsx", "range": rng})

    def test_excel_external_formula_and_duplicate_sheet_rejected(self):
        for f in ('=WEBSERVICE("https://example.com")', "='[private.xlsx]Sheet1'!A1", '=cmd|"x"!A1'):
            with self.subTest(formula=f), self.assertRaises(ToolError):
                office.cell_value({"formula": f})
        s = self.spec({"sheets": [{"name": "Data", "rows": [[1]]}, {"name": "data", "rows": [[2]]}]})
        with self.assertRaises(ToolError):
            office.compose({"format": "xlsx", "spec": s, "output": "duplicate.xlsx"})
        self.assertFalse((self.root/"duplicate.xlsx").exists())

    def test_powerpoint_editable_slides_notes_and_edit(self):
        s = self.spec({"slides": [{"title": "Xin chào {{name}}", "bullets": ["Tài liệu có dấu"], "notes": "{{name}} trình bày"}]})
        office.compose({"format": "pptx", "spec": s, "output": "slides.pptx"})
        s = self.spec({"replacements": [{"find": "{{name}}", "replace": "Ngọc Anh", "expected": 2}]})
        office.edit({"path": "slides.pptx", "spec": s, "output": "edited.pptx"})
        slide = office.inspect({"path": "edited.pptx"})["slides"][0]
        self.assertIn("Ngọc Anh", slide["text"][0])
        self.assertIn("Ngọc Anh", slide["notes"])

    def test_pdf_create_extract_merge_select_rotate_and_render_actual_pixels(self):
        from pypdf import PdfReader
        from PIL import Image
        p = self.pdf()
        actual = office.inspect({"path": p})
        self.assertIn("Báo cáo tiếng Việt", actual["pages"][0]["text"])
        office.pdf_modify({"paths": [p,p], "output": "merged.pdf", "rotation": 90})
        doc = PdfReader(self.root / "merged.pdf")
        self.assertEqual(len(doc.pages), 2)
        self.assertEqual(doc.pages[0].rotation, 90)
        office.pdf_modify({"paths": ["merged.pdf"], "pages": [1], "output": "selected.pdf"})
        self.assertEqual(len(PdfReader(self.root / "selected.pdf").pages), 1)
        office.render({"path": "selected.pdf", "output": "page.png"})
        with Image.open(self.root / "page.png") as im:
            self.assertGreater(im.width, im.height)
            self.assertGreater(len(im.convert("L").getcolors(im.width*im.height)), 2)

    def test_pdf_form_fields_are_filled_and_unknown_fields_rejected(self):
        from reportlab.pdfgen import canvas
        from pypdf import PdfReader
        c = canvas.Canvas(str(self.root / "form.pdf"))
        c.acroForm.textfield(name="customer", x=50, y=700, width=200, height=30)
        c.showPage(); c.save()
        office.pdf_modify({"paths": ["form.pdf"], "fields": {"customer": "Ngoc Anh"}, "output": "filled.pdf"})
        self.assertEqual(PdfReader(self.root/"filled.pdf").get_fields()["customer"]["/V"], "Ngoc Anh")
        with self.assertRaises(ToolError):
            office.pdf_modify({"paths": ["form.pdf"], "fields": {"unknown": "x"}, "output": "fail.pdf"})

    def test_office_zip_expansion_guard(self):
        import zipfile
        with zipfile.ZipFile(self.root / "bomb.docx", "w", compression=zipfile.ZIP_DEFLATED) as z:
            z.writestr("word/document.xml", b"A"*(2*1024*1024))
        with self.assertRaises(ToolError):
            office.inspect({"path": "bomb.docx"})

    def test_ocr_native_pdf_route_does_not_contact_model(self):
        p = self.pdf()
        self.assertEqual(ocr.assess({"path": p})["pages"][0]["route"], "native-text")
        with patch.object(ocr, "http", side_effect=AssertionError("must not call model")):
            result = ocr.recognize({"path": p, "output": "extract.md"})
        self.assertEqual(result["pages"][0]["engine"], "native-text")
        self.assertIn("Báo cáo tiếng Việt", (self.root/"extract.md").read_text())
        self.assertIn("source-page: 1", (self.root/"extract.md").read_text())

    def test_ocr_image_data_url_actual_http_and_force_pdf(self):
        from PIL import Image
        Image.new("RGB", (200, 100), "white").save(self.root / "scan.png")
        payload = {"choices": [{"finish_reason": "stop", "message": {"content": "Đọc đúng số: 125"}}]}
        with service(lambda r: (200, json.dumps(payload).encode(), {})) as (url, calls):
            with patch.dict(os.environ, {"OCR_BINDING": json.dumps({"base_url": url+"/v1", "model": "deepseek-ai/DeepSeek-OCR-2"})}):
                result = ocr.recognize({"path": "scan.png", "output": "scan.md"})
        self.assertEqual(result["pages"][0]["engine"], "deepseek-ocr")
        request = json.loads(calls[0][2])
        self.assertEqual(calls[0][1], "/v1/chat/completions")
        self.assertEqual(request["messages"][0]["content"][0]["text"], "<|grounding|>Convert the document to markdown.")
        self.assertNotIn("<image>", request["messages"][0]["content"][0]["text"])
        data_url = request["messages"][0]["content"][1]["image_url"]["url"]
        self.assertTrue(base64.b64decode(data_url.split(",",1)[1]).startswith(b"\x89PNG"))
        pdf = self.pdf()
        with service(lambda r: (200, json.dumps(payload).encode(), {})) as (url, calls):
            with patch.dict(os.environ, {"OCR_BINDING": json.dumps({"base_url": url, "model": "deepseek-ai/DeepSeek-OCR-2"})}):
                result = ocr.recognize({"path": pdf, "force": True, "output": "forced.md"})
        self.assertEqual(result["pages"][0]["engine"], "deepseek-ocr")
        encoded = json.loads(calls[0][2])["messages"][0]["content"][1]["image_url"]["url"]
        with Image.open(io.BytesIO(base64.b64decode(encoded.split(",", 1)[1]))) as image:
            self.assertGreater(image.width, 100)
            self.assertGreater(image.height, image.width)

    def test_ocr_truncated_inference_not_published_and_generic_chat_rejected(self):
        from PIL import Image
        Image.new("RGB", (20, 20)).save(self.root/"scan.png")
        payload = {"choices": [{"finish_reason": "length", "message": {"content": "partial"}}]}
        with service(lambda r: (200, json.dumps(payload).encode(), {})) as (url, calls):
            for model in ("deepseek-ai/DeepSeek-OCR-2", "deepseek-chat"):
                with self.subTest(model=model), patch.dict(os.environ, {"OCR_BINDING": json.dumps({"base_url": url, "model": model})}), self.assertRaises(ToolError):
                    ocr.recognize({"path": "scan.png", "output": "fail.md"})
        self.assertFalse((self.root/"fail.md").exists())
        self.assertEqual(len(calls), 1)

    def test_ocr_duplicate_and_excess_page_selection_rejected(self):
        p = self.pdf()
        for pages in ([0,0], list(range(9)), [-1], [True]):
            with self.subTest(pages=pages), self.assertRaises(ToolError):
                ocr.assess({"path": p, "pages": pages})

    def test_http_redirect_never_forwards_bearer_token(self):
        with service(lambda r: (302, b"", {"Location": "/stolen"})) as (url, calls):
            with self.assertRaises(ToolError):
                tts.http(url + "/source", token="synthetic-test-token")
        self.assertEqual(len(calls), 1)

    def test_http_mutation_oversized_success_is_unknown_not_safe_to_retry(self):
        with service(lambda r: (200, b"x"*1025, {})) as (url, calls):
            with self.assertRaises(Uncertain):
                google.http(url, method="POST", data={}, maximum=1024, mutation=True)
        self.assertEqual(len(calls), 1)

    def test_http_408_mutation_is_unknown_and_never_replayed(self):
        with service(lambda r: (408, b"private vendor diagnostic", {})) as (url, calls):
            with self.assertRaises(Uncertain) as result:
                google.http(url, method="POST", data={}, mutation=True)
            self.assertNotIn("private", str(result.exception))
        self.assertEqual(len(calls), 1)

    def test_public_result_size_is_bounded_without_partial_stdout(self):
        from docx import Document
        doc = Document()
        for _ in range(100):
            doc.add_paragraph("x"*10000)
        doc.save(self.root / "large.docx")
        call = {"api_version": "soot/v1", "hook": "tool.call", "soot": "test",
                "operation": "office_inspect", "input": {"path": "large.docx", "limit": 100}}
        result = subprocess.run([str(REPO / "packs/office/tools/run"), "inspect"],
                                input=json.dumps(call), text=True, capture_output=True,
                                env={"FILE_ROOT": str(self.root), "PYTHON_EXECUTABLE": sys.executable})
        self.assertEqual(result.returncode, 1)
        self.assertEqual(result.stdout, "")
        self.assertIn("reduce", result.stderr)

    def test_tts_real_http_client_pcm_duration_transcript_and_voice_guard(self):
        with service(lambda r: (200, wav_bytes(.25), {"Content-Type": "audio/wav"})) as (url, calls):
            cfg = {"engine": "omnivoice", "base_url": url, "voices": ["reference"]}
            with patch.dict(os.environ, {"TTS_BINDING": json.dumps(cfg)}):
                result = tts.synthesize({"text": "Xin chào Việt Nam", "voice": "reference", "output": "speech.wav"})
                with self.assertRaises(ToolError):
                    tts.synthesize({"text": "x", "voice": "other", "output": "denied.wav"})
        self.assertEqual(result["duration"], .25)
        self.assertEqual(result["transcript"], "Xin chào Việt Nam")
        self.assertEqual(result["text_sha256"], hashlib.sha256("Xin chào Việt Nam".encode()).hexdigest())
        self.assertEqual(len(calls), 1)
        self.assertEqual(json.loads(calls[0][2])["num_steps"], 32)

    def test_tts_invalid_or_truncated_audio_not_published(self):
        multichannel = io.BytesIO()
        with wave.open(multichannel, "wb") as wav:
            wav.setnchannels(3); wav.setsampwidth(2); wav.setframerate(24000); wav.writeframes(b"\0"*600)
        for content in (b"<html>error</html>", wav_bytes()[:-100], multichannel.getvalue()):
            with self.subTest(content=content[:8]), service(lambda r: (200, content, {})) as (url, calls):
                with patch.dict(os.environ, {"TTS_BINDING": json.dumps({"engine": "omnivoice", "base_url": url, "voices": ["reference"]})}):
                    with self.assertRaises(Exception):
                        tts.synthesize({"text": "Xin chào", "voice": "reference", "output": "fail.wav"})
                self.assertFalse((self.root/"fail.wav").exists())

    def test_tts_existing_output_never_spends_another_inference(self):
        (self.root / "existing.wav").write_bytes(wav_bytes())
        with service(lambda r: (200, wav_bytes(), {})) as (url, calls):
            with patch.dict(os.environ, {"TTS_BINDING": json.dumps({"engine":"omnivoice", "base_url":url, "voices":["reference"]})}):
                with self.assertRaises(ToolError):
                    tts.synthesize({"text":"Xin chào", "voice":"reference", "output":"existing.wav"})
        self.assertEqual(calls, [])

    def test_tts_local_inspection_detects_declared_missing_samples(self):
        (self.root / "truncated.wav").write_bytes(wav_bytes()[:-20])
        with self.assertRaises(ToolError):
            tts.inspect({"path": "truncated.wav"})

    def google_token(self):
        return patch.dict(os.environ, {"GOOGLE_WORKSPACE_BINDING": json.dumps({"account_id": "test@example.invalid", "access_token": "synthetic-test-token", "expires_at": time.time()+3600})})

    def test_google_tabs_nested_content_utf16_and_pagination(self):
        data = {"title": "Tài liệu", "revisionId": "rev-actual", "tabs": [{"tabProperties": {"tabId": "t1", "title": "One"},
                    "childTabs": [{"tabProperties": {"tabId": "t2", "title": "Two"}, "documentTab": {"body": {"content": [
                        {"startIndex": 1, "endIndex": 6, "paragraph": {"elements": [{"textRun": {"content": "😀 Việt"}}]}}]}}}]}]}
        with self.google_token(), patch.object(google, "http", return_value=(json.dumps(data).encode(), {})) as transport:
            result = google.docs_read({"document_id": "sample_document_123", "tab_id": "t2", "limit": 1})
        self.assertEqual(result["content"][0]["text"], "😀 Việt")
        self.assertEqual(result["revision_id"], "rev-actual")
        self.assertEqual(len(result["tabs"]), 2)
        self.assertIn("includeTabsContent=true", transport.call_args.args[0])
        self.assertIn("UTF-16", result["index_unit"])

    def test_google_expired_credentials_fail_before_transport(self):
        with patch.dict(os.environ, {"GOOGLE_WORKSPACE_BINDING": json.dumps({"access_token": "synthetic", "expires_at": 1})}), patch.object(google, "http") as transport:
            with self.assertRaises(ToolError):
                google.docs_read({"document_id": "sample_document_123"})
        transport.assert_not_called()

    def test_google_create_durable_duplicate_and_request_conflict(self):
        response = {"documentId": "new_document_123", "title": "Title"}
        with self.google_token(), patch.object(google, "http", return_value=(json.dumps(response).encode(), {})) as transport:
            first = google.docs_create({"title": "Title", "action_id": "create-1"})
            second = google.docs_create({"title": "Title", "action_id": "create-1"})
            with self.assertRaises(ToolError):
                google.docs_create({"title": "Other", "action_id": "create-1"})
        self.assertEqual(first["document_id"], second["document_id"])
        self.assertTrue(second["replayed_receipt"])
        self.assertEqual(transport.call_count, 1)

    def test_google_unknown_write_never_resends_on_next_process_state(self):
        with self.google_token(), patch.object(google, "http", side_effect=Uncertain("synthetic disconnect")) as transport:
            with self.assertRaises(Uncertain):
                google.docs_create({"title": "Title", "action_id": "uncertain-1"})
        # Fresh call, persisted pending state, different transport: it must not execute.
        with self.google_token(), patch.object(google, "http") as transport:
            with self.assertRaises(Uncertain):
                google.docs_create({"title": "Title", "action_id": "uncertain-1"})
        transport.assert_not_called()

    def test_google_receipts_are_account_scoped_and_tokens_can_rotate(self):
        response = {"documentId":"new_document_123", "title":"Title"}
        cfg = {"account_id":"first@example.invalid", "access_token":"synthetic-1", "expires_at":time.time()+3600}
        with patch.object(google, "http", return_value=(json.dumps(response).encode(), {})) as transport:
            with patch.dict(os.environ, {"GOOGLE_WORKSPACE_BINDING":json.dumps(cfg)}):
                google.docs_create({"title":"Title", "action_id":"same-action"})
            cfg["access_token"] = "synthetic-rotated"
            with patch.dict(os.environ, {"GOOGLE_WORKSPACE_BINDING":json.dumps(cfg)}):
                self.assertTrue(google.docs_create({"title":"Title", "action_id":"same-action"})["replayed_receipt"])
            cfg["account_id"] = "second@example.invalid"
            with patch.dict(os.environ, {"GOOGLE_WORKSPACE_BINDING":json.dumps(cfg)}):
                self.assertNotIn("replayed_receipt", google.docs_create({"title":"Title", "action_id":"same-action"}))
        self.assertEqual(transport.call_count, 2)

    def test_google_wrong_success_identity_remains_pending_and_never_resends(self):
        s = self.spec({"requests":[{"insertText":{"endOfSegmentLocation":{"tabId":"t1"}, "text":"Xin chào"}}]})
        args = {"document_id":"sample_document_123", "revision_id":"current", "spec":s, "action_id":"wrong-success"}
        with self.google_token(), patch.object(google, "http", return_value=(b'{"documentId":"wrong_document_123","replies":[]}', {})) as transport:
            with self.assertRaises(Uncertain):
                google.docs_batch_update(args)
            with self.assertRaises(Uncertain):
                google.docs_batch_update(args)
        self.assertEqual(transport.call_count, 1)

    def test_google_batch_requires_revision_precise_mask_and_tab(self):
        req = {"insertText": {"endOfSegmentLocation": {"tabId": "t1"}, "text": "Xin chào\n"}}
        s = self.spec({"requests": [req]})
        with self.google_token(), patch.object(google, "http", return_value=(b'{"documentId":"sample_document_123","replies":[{}],"writeControl":{"requiredRevisionId":"next"}}', {})) as transport:
            google.docs_batch_update({"document_id": "sample_document_123", "revision_id": "current", "spec": s, "action_id": "edit-1"})
        self.assertEqual(transport.call_args.kwargs["data"]["writeControl"], {"requiredRevisionId": "current"})
        for bad in ({"insertText": {"text": "x", "location": {"index": 1}}},
                    {"updateTextStyle": {"range": {"tabId": "t1", "startIndex": 1}, "fields": "*"}},
                    {"replaceAllText": {"containsText": {"text": "x"}, "replaceText": "y"}},
                    {"deleteNamedRange": {"name": "x"}}):
            s = self.spec({"requests": [bad]})
            with self.subTest(bad=bad), self.assertRaises(ToolError):
                google.docs_batch_update({"document_id": "sample_document_123", "revision_id": "current", "spec": s, "action_id": "bad"})

    def test_google_export_rejects_html_and_is_bounded_to_drive_endpoint(self):
        with self.google_token(), patch.object(google, "http", return_value=(b"<html>login</html>", {})) as transport:
            with self.assertRaises(ToolError):
                google.docs_export({"document_id": "sample_document_123", "format": "pdf", "output": "bad.pdf"})
        self.assertTrue(transport.call_args.args[0].startswith("https://www.googleapis.com/drive/v3/files/"))
        self.assertFalse((self.root/"bad.pdf").exists())

    def video_config(self):
        ffmpeg, ffprobe = shutil.which("ffmpeg"), shutil.which("ffprobe")
        if not ffmpeg or not ffprobe:
            raise RuntimeError("FFmpeg and FFprobe are required, not skipped")
        return patch.dict(os.environ, {"VIDEO_BINDING": json.dumps({"ffmpeg": ffmpeg, "ffprobe": ffprobe})})

    def test_sheets_bounded_formula_read_preserves_unicode_and_trailing_blank_snapshot(self):
        payload = {"range":"'Chi phí'!A1:B2", "values":[["Mục","=SUM(B2:B3)"],[125]]}
        with self.google_token(), patch.object(google,"http",return_value=(json.dumps(payload).encode(),{})) as transport:
            result = google.sheets_read({"spreadsheet_id":"spreadsheet_123", "range":"'Chi phí'!A1:B2"})
        self.assertEqual(result["values"][0][1],"=SUM(B2:B3)")
        self.assertEqual(result["render"],"FORMULA")
        self.assertEqual(len(result["snapshot_sha256"]),64)
        self.assertIn("sheets.googleapis.com/v4/spreadsheets/",transport.call_args.args[0])
        self.assertIn("%27Chi%20ph%C3%AD%27%21A1%3AB2",transport.call_args.args[0])
        blank_explicit = {"values":[["Mục","=SUM(B2:B3)"],[125,""]]}
        with self.google_token(),patch.object(google,"http",return_value=(json.dumps(blank_explicit).encode(),{})):
            second = google.sheets_read({"spreadsheet_id":"spreadsheet_123","range":"'Chi phí'!A1:B2"})
        self.assertEqual(second["snapshot_sha256"],result["snapshot_sha256"])
        for value in ('A:A','Sheet1','Sheet1!A1:ZZZ9999','Sheet1!A0:B2','Sheet1!B2:A1'):
            with self.subTest(range=value), self.assertRaises(ToolError):google.a1(value)

    def test_sheets_exact_update_preflight_raw_literal_and_duplicate_receipt(self):
        read={"values":[["old",125]]}
        update={"spreadsheetId":"spreadsheet_123","updatedRange":"Sheet1!A1:B1","updatedRows":1,"updatedCells":2}
        with self.google_token(),patch.object(google,"http",return_value=(json.dumps(read).encode(),{})):
            snapshot=google.sheets_read({"spreadsheet_id":"spreadsheet_123","range":"Sheet1!A1:B1"})["snapshot_sha256"]
        spec=self.spec({"values":[["=literal",250]]})
        args={"spreadsheet_id":"spreadsheet_123","range":"Sheet1!A1:B1","spec":spec,"expected_snapshot":snapshot,"action_id":"update-sheet"}
        with self.google_token(),patch.object(google,"http",side_effect=[(json.dumps(read).encode(),{}),(json.dumps(update).encode(),{})]) as transport:
            result=google.sheet_write(args)
            duplicate=google.sheet_write(args)
        self.assertEqual(transport.call_count,2)
        self.assertTrue(duplicate["replayed_receipt"])
        self.assertIn("valueInputOption=RAW",transport.call_args.args[0])
        self.assertEqual(transport.call_args.kwargs["method"],"PUT")
        self.assertEqual(transport.call_args.kwargs["data"]["values"][0][0],"=literal")
        self.assertIn("non-atomic",result["concurrency"])

    def test_sheets_stale_snapshot_never_writes_and_unknown_append_never_replays(self):
        spec=self.spec({"values":[["x"]]})
        args={"spreadsheet_id":"spreadsheet_123","range":"Sheet1!A1","spec":spec,"expected_snapshot":"0"*64,"action_id":"stale-sheet"}
        with self.google_token(),patch.object(google,"http",return_value=(b'{"values":[["changed"]]}',{})) as transport:
            with self.assertRaises(ToolError):google.sheet_write(args)
        self.assertEqual(transport.call_count,1)
        self.assertEqual(transport.call_args.kwargs["method"],"GET")
        args.pop("expected_snapshot");args["action_id"]="append-unknown"
        with self.google_token(),patch.object(google,"http",side_effect=Uncertain("disconnect")) as transport:
            with self.assertRaises(Uncertain):google.sheet_write(args,True)
            with self.assertRaises(Uncertain):google.sheet_write(args,True)
        self.assertEqual(transport.call_count,1)

    def test_sheets_create_append_actual_range_and_input_validation(self):
        create={"spreadsheetId":"spreadsheet_123"}
        with self.google_token(),patch.object(google,"http",return_value=(json.dumps(create).encode(),{})) as transport:
            google.sheets_create({"title":"Ngân sách","tabs":["Chi phí","Tổng"],"action_id":"sheet-create"})
        self.assertEqual(transport.call_args.kwargs["data"]["sheets"][1]["properties"]["title"],"Tổng")
        spec=self.spec({"values":[["=SUM(B2:B3)"]],"input_option":"USER_ENTERED"})
        response={"spreadsheetId":"spreadsheet_123","updates":{"updatedRange":"Sheet1!A21","updatedRows":1}}
        with self.google_token(),patch.object(google,"http",return_value=(json.dumps(response).encode(),{})) as transport:
            result=google.sheet_write({"spreadsheet_id":"spreadsheet_123","range":"Sheet1!A1:A20","spec":spec,"action_id":"sheet-append"},True)
        self.assertEqual(result["updates"]["updatedRange"],"Sheet1!A21")
        self.assertIn("insertDataOption=INSERT_ROWS",transport.call_args.args[0])
        for value in ([[None]],[[float('nan')]],[["=IMPORTXML(\"https://example.invalid\")"]]):
            with self.subTest(value=value),self.assertRaises(ToolError):google.matrix(value,1,1,input_option="USER_ENTERED")

    def test_drive_listing_is_paginated_scoped_and_query_is_not_model_input(self):
        with self.google_token(),patch.object(google,"http",return_value=(b'{"files":[],"nextPageToken":"next"}',{})) as transport:
            result=google.drive_list({"parent_id":"folder_resource_123","kind":"sheets","limit":5,"page_token":"opaque"})
        from urllib.parse import urlsplit,parse_qs
        query=parse_qs(urlsplit(transport.call_args.args[0]).query)
        self.assertEqual(query["pageSize"],["5"])
        self.assertIn("'folder_resource_123' in parents",query['q'][0])
        self.assertEqual(result["nextPageToken"],"next")
        with self.assertRaises(ToolError):google.drive_list({"q":"name contains 'secret'"})

    def test_drive_binary_download_verifies_checksum_and_export_routes_native_sheets(self):
        raw=b'file bytes'
        meta={"id":"file_resource_123","mimeType":"text/plain","size":str(len(raw)),"md5Checksum":hashlib.md5(raw).hexdigest(),"version":"12"}
        with self.google_token(),patch.object(google,"http",side_effect=[(json.dumps(meta).encode(),{}),(raw,{})]):
            google.drive_download({"file_id":"file_resource_123","output":"download.txt"})
        self.assertEqual((self.root/'download.txt').read_bytes(),raw)
        with self.google_token(),patch.object(google,"http",side_effect=[(json.dumps(meta).encode(),{}),(b'wrongbytes',{})]):
            with self.assertRaises(ToolError):google.drive_download({"file_id":"file_resource_123","output":"corrupt.txt"})
        self.assertFalse((self.root/'corrupt.txt').exists())
        from openpyxl import Workbook
        workbook=Workbook();workbook.active["A1"]="Báo cáo";stream=io.BytesIO();workbook.save(stream);workbook.close()
        meta={"id":"spreadsheet_123","mimeType":"application/vnd.google-apps.spreadsheet"}
        with self.google_token(),patch.object(google,"http",side_effect=[(json.dumps(meta).encode(),{}),(stream.getvalue(),{})]) as transport:
            google.drive_export({"file_id":"spreadsheet_123","format":"xlsx","output":"export.xlsx"})
        self.assertIn("/export?",transport.call_args.args[0])
        self.assertEqual(office.inspect({"path":"export.xlsx","range":"A1:A1"})["rows"],[["Báo cáo"]])

    def test_drive_multipart_upload_preserves_bytes_metadata_and_deduplicates(self):
        raw='Nội dung chính xác\r\n'.encode();(self.root/'source.txt').write_bytes(raw)
        response={"id":"file_resource_123","name":"Báo cáo.txt"}
        args={"path":"source.txt","name":"Báo cáo.txt","parent_id":"folder_resource_123","action_id":"upload-one"}
        with self.google_token(),patch.object(google,"http",return_value=(json.dumps(response).encode(),{})) as transport:
            result=google.drive_upload(args);duplicate=google.drive_upload(args)
        self.assertEqual(transport.call_count,1)
        self.assertTrue(duplicate["replayed_receipt"])
        from email.parser import BytesParser
        from email.policy import default
        mime=BytesParser(policy=default).parsebytes(('Content-Type: '+transport.call_args.kwargs['headers']['Content-Type']+'\r\nMIME-Version: 1.0\r\n\r\n').encode()+transport.call_args.kwargs['data'])
        parts=list(mime.iter_parts())
        metadata=json.loads(parts[0].get_payload(decode=True))
        self.assertEqual(metadata['name'],'Báo cáo.txt')
        self.assertEqual(metadata['parents'],['folder_resource_123'])
        self.assertEqual(parts[1].get_payload(decode=True),raw)
        self.assertEqual(result['source_sha256'],hashlib.sha256(raw).hexdigest())

    def test_drive_invalid_office_export_is_not_published(self):
        import zipfile
        meta={"id":"spreadsheet_123","mimeType":"application/vnd.google-apps.spreadsheet"}
        empty=io.BytesIO()
        with zipfile.ZipFile(empty,"w") as archive:
            archive.writestr('unrelated.txt','Not a workbook')
        for index,raw in enumerate((b'PK-not-a-valid-zip',empty.getvalue())):
            filename=f'invalid-{index}.xlsx'
            with self.google_token(),patch.object(google,"http",side_effect=[(json.dumps(meta).encode(),{}),(raw,{})]):
                with self.assertRaises(ToolError):google.drive_export({"file_id":"spreadsheet_123","format":"xlsx","output":filename})
            self.assertFalse((self.root/filename).exists())

    def test_drive_missing_version_and_oversized_upload_never_contact_google(self):
        with self.google_token(),patch.object(google,"http") as transport:
            with self.assertRaises(ToolError):
                google.drive_update({"file_id":"file_resource_123","expected_version":"None","name":"x","action_id":"bad-version"})
            with (self.root/'large.txt').open('wb') as stream:
                stream.truncate(5*1024*1024+1)
            with self.assertRaises(ToolError):google.drive_upload({"path":"large.txt","action_id":"too-large"})
        transport.assert_not_called()

    def test_drive_version_preflight_folder_creation_and_parent_move(self):
        meta={"id":"file_resource_123","version":"12","parents":["old_parent_123"]}
        response={"id":"file_resource_123","version":"13","parents":["new_parent_123"]}
        args={"file_id":"file_resource_123","expected_version":"12","parent_id":"new_parent_123","name":"Renamed","action_id":"move-one"}
        with self.google_token(),patch.object(google,"http",side_effect=[(json.dumps(meta).encode(),{}),(json.dumps(response).encode(),{})]) as transport:
            result=google.drive_update(args)
        self.assertEqual(transport.call_args.kwargs['method'],'PATCH')
        self.assertIn('addParents=new_parent_123',transport.call_args.args[0])
        self.assertIn('removeParents=old_parent_123',transport.call_args.args[0])
        self.assertTrue(result['parent_permissions_may_apply'])
        args['action_id']='stale-move';args['expected_version']='11'
        with self.google_token(),patch.object(google,"http",return_value=(json.dumps(meta).encode(),{})) as transport:
            with self.assertRaises(ToolError):google.drive_update(args)
        self.assertEqual(transport.call_count,1)
        with self.google_token(),patch.object(google,"http",return_value=(b'{"id":"folder_resource_123","mimeType":"application/vnd.google-apps.folder"}',{})):
            result=google.drive_folder({"name":"Reports","parent_id":"root","action_id":"folder-one"})
        self.assertFalse(result['sharing_changed'])

    def test_workspace_public_wrapper_uses_new_identity_and_rejects_expired_auth(self):
        call={"api_version":"soot/v1","hook":"tool.call","soot":"test","operation":"google-workspace_sheets_read","input":{"spreadsheet_id":"spreadsheet_123","range":"Sheet1!A1"}}
        result=subprocess.run([str(REPO/'packs/google-workspace/tools/run'),'sheets_read'],input=json.dumps(call),text=True,capture_output=True,
                              env={"FILE_ROOT":str(self.root),"PYTHON_EXECUTABLE":sys.executable,"GOOGLE_WORKSPACE_BINDING":'{"access_token":"synthetic","expires_at":1}'})
        self.assertEqual(result.returncode,1)
        self.assertIn('expired',result.stderr)
        call['operation']='google-doc_read'
        result=subprocess.run([str(REPO/'packs/google-workspace/tools/run'),'sheets_read'],input=json.dumps(call),text=True,capture_output=True,
                              env={"FILE_ROOT":str(self.root),"PYTHON_EXECUTABLE":sys.executable})
        self.assertIn('operation mismatch',result.stderr)

    def video_spec(self):
        from PIL import Image, ImageDraw
        image = Image.new("RGB", (360,640), "navy")
        ImageDraw.Draw(image).rectangle((80,100,280,400), fill="gold")
        image.save(self.root/"a.png")
        (self.root/"voice.wav").write_bytes(wav_bytes(.8))
        return {"width":360,"height":640,"fps":24,"clips":[{"path":"a.png","duration":.6,"zoom_end":1.08},
                {"path":"a.png","duration":.6,"fade_out":.1}],
                "audio":[{"path":"voice.wav","role":"voice","start":.1,"gain":.8}],
                "font":self.font(),"caption_size":24,
                "captions":[{"start":.1,"end":1.0,"text":"Xin chào Việt Nam","timing":"manual"}]}

    def test_video_actual_encode_audio_caption_decode_contact_sheet(self):
        spec = self.spec(self.video_spec())
        with self.video_config():
            validation = video.validate({"spec":spec})
            self.assertTrue(validation["valid"])
            result = video.render({"spec":spec,"output":"out.mp4"})
            self.assertEqual(result["video"][0]["width"],360)
            self.assertEqual(result["video"][0]["height"],640)
            self.assertEqual(result["audio"][0]["codec_name"],"aac")
            sheet = video.contact_sheet({"path":"out.mp4","output":"sheet.png","times":[.1,.6,1.0]})
            self.assertEqual(sheet["samples"],3)
            self.assertTrue(sheet["all_frames_decoded"])

    def test_video_bad_caption_bounds_directives_and_fake_alignment_rejected(self):
        spec = self.video_spec()
        for change in ({"text":"{\\pos(1,1)}hijack"},{"timing":"estimated"},
                       {"text":"X"*100},{"end":5}):
            candidate = json.loads(json.dumps(spec));candidate["captions"][0].update(change)
            s = self.spec(candidate)
            with self.subTest(change=change), self.video_config(), self.assertRaises(ToolError):
                video.validate({"spec":s})

    def test_video_measured_source_ranges_and_no_overwrite(self):
        s = self.spec(self.video_spec())
        with self.video_config():
            video.render({"spec":s,"output":"source.mp4"})
            with self.assertRaises(ToolError):
                video.render({"spec":s,"output":"source.mp4"})
            bad = self.spec({"width":360,"height":640,"clips":[{"path":"source.mp4","source_start":1,"duration":2}]})
            with self.assertRaises(ToolError):
                video.validate({"spec":bad})

    def test_hyperframes_qa_findings_survive_expected_nonzero_exit(self):
        (self.root/'motion').mkdir()
        (self.root/'motion/index.html').write_text('<div>fixture</div>')
        executable = str(Path(sys.executable).resolve())
        cli = self.root/'fixture-hyperframes.py'
        cfg = {"node":executable, "cli":str(cli)}
        report = {"ok":False, "errorCount":1, "findings":[{"code":"multiple_root_compositions", "severity":"error"}]}
        cli.write_text('import sys\nassert sys.argv[-1] == "--json"\nprint('+repr(json.dumps(report))+')\nraise SystemExit(1)\n')
        with patch.dict(os.environ, {"HYPERFRAMES_BINDING":json.dumps(cfg)}):
            result = video.hyperframes({"project":"motion", "action":"check"})
        self.assertFalse(result["passed"])
        self.assertEqual(result["report"]["findings"][0]["code"], "multiple_root_compositions")
        cli.write_text('print("not an audit")\nraise SystemExit(1)\n')
        with patch.dict(os.environ, {"HYPERFRAMES_BINDING":json.dumps(cfg)}):
            with self.assertRaises(ToolError):
                video.hyperframes({"project":"motion", "action":"check"})

    def test_video_commands_bound_output_and_timeout_including_descendants(self):
        with self.assertRaises(ToolError):
            video.command([sys.executable,'-c','import sys; sys.stdout.buffer.write(b"x"*(2*1024*1024+1))'])
        marker = self.root/'child-survived.txt'
        child = 'import time;from pathlib import Path;time.sleep(.5);Path('+repr(str(marker))+').write_text("survived")'
        parent = 'import subprocess,sys;subprocess.Popen([sys.executable,"-c",'+repr(child)+'])'
        with self.assertRaises(ToolError):
            video.command([sys.executable,'-c',parent],timeout=.1)
        time.sleep(.6)
        self.assertFalse(marker.exists(), 'timed-out CLI left a child process running')
        self.assertEqual(video.command([sys.executable,'-c','print("known finding");raise SystemExit(1)'],allowed_codes=(0,1)),b'known finding\n')

    def test_video_repeated_assets_are_probed_once_per_timeline(self):
        self.video_spec()
        s = self.spec({'width':360,'height':640,'clips':[{'path':'a.png','duration':1} for _ in range(10)],
                       'audio':[{'path':'voice.wav','role':'voice','start':0}, {'path':'voice.wav','role':'sfx','start':1}]})
        with self.video_config(), patch.object(video,'probe_file',wraps=video.probe_file) as measured:
            data = video.timeline({'spec':s})
        self.assertEqual(len(data['clips']),10)
        self.assertEqual(measured.call_count,2)

    def test_video_invalid_probe_metadata_is_rejected_before_decode(self):
        for data in [{'format':{'duration':'NaN'},'streams':[]},
                     {'format':{'duration':'1'},'streams':[{'codec_type':'video','width':32768,'height':32768}]}]:
            with patch.object(video,'command',return_value=json.dumps(data).encode()), self.assertRaises(ToolError):
                video.probe_file(self.root/'source.mp4',{'ffprobe':'fixture'})

    def test_tts_silent_pcm_is_identified_and_never_published_as_speech(self):
        silent = io.BytesIO()
        with wave.open(silent,'wb') as f:
            f.setnchannels(1);f.setsampwidth(2);f.setframerate(24000);f.writeframes(b'\0'*4800)
        (self.root/'silent.wav').write_bytes(silent.getvalue())
        self.assertTrue(tts.inspect({'path':'silent.wav'})['digital_silence'])
        cfg={'base_url':'http://127.0.0.1:8042','engine':'omnivoice','voices':['fixture']}
        with patch.dict(os.environ,{'TTS_BINDING':json.dumps(cfg)}), patch.object(tts,'http',return_value=(silent.getvalue(),{})):
            with self.assertRaises(ToolError):
                tts.synthesize({'voice':'fixture','text':'Có tiếng nói','output':'speech.wav'})
        self.assertFalse((self.root/'speech.wav').exists())

    def test_word_table_text_is_never_silently_truncated(self):
        for value in ['x'*10001, {'unexpected':'object'}, float('inf')]:
            s = self.spec({'blocks':[{'kind':'table','rows':[[value]]}]})
            with self.assertRaises(ToolError):
                office.compose({'format':'docx','spec':s,'output':'bad.docx'})
            self.assertFalse((self.root/'bad.docx').exists())

    def test_powerpoint_large_table_remains_readable_with_bounded_preview(self):
        from pptx import Presentation
        from pptx.util import Inches
        doc = Presentation(); slide=doc.slides.add_slide(doc.slide_layouts[6]); table=slide.shapes.add_table(31,21, Inches(.1), Inches(.1), Inches(9), Inches(6)).table
        table.cell(30,20).text='outside preview'
        doc.save(self.root/'large-table.pptx')
        data = office.inspect({'path':'large-table.pptx'})
        self.assertEqual(len(data['slides'][0]['tables'][0]),30)
        self.assertEqual(len(data['slides'][0]['tables'][0][0]),20)
        self.assertEqual(data['slides'][0]['table_dimensions'][0],{'rows':31,'columns':21,'truncated':True})

    def test_powerpoint_grouped_text_can_be_read_and_edited(self):
        from pptx import Presentation
        from pptx.util import Inches
        doc=Presentation();slide=doc.slides.add_slide(doc.slide_layouts[6]);group=slide.shapes.add_group_shape()
        group.shapes.add_textbox(Inches(1),Inches(1),Inches(4),Inches(2)).text='Nhóm {{name}}'
        doc.save(self.root/'grouped.pptx')
        self.assertEqual(office.inspect({'path':'grouped.pptx'})['slides'][0]['text'],['Nhóm {{name}}'])
        s=self.spec({'replacements':[{'find':'{{name}}','replace':'Soot','expected':1}]})
        office.edit({'path':'grouped.pptx','spec':s,'output':'edited-group.pptx'})
        self.assertEqual(office.inspect({'path':'edited-group.pptx'})['slides'][0]['text'],['Nhóm Soot'])

    def test_google_nonfinite_token_expiry_is_rejected(self):
        for expires in [float('inf'),float('nan')]:
            with patch.dict(os.environ,{'GOOGLE_WORKSPACE_BINDING':json.dumps({'access_token':'fixture','expires_at':expires})}), self.assertRaises(ToolError):
                google.auth()


if __name__ == "__main__":
    unittest.main(verbosity=2)
