"""Check live-provider artifacts independently of the pack's own replies."""
import hashlib
import json
from pathlib import Path
import subprocess
import sys
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader
from PIL import Image

root = Path(sys.argv[1])
doc = Document(root/'result.docx')
assert doc.paragraphs[0].text == 'Báo cáo tiếng Việt'
assert doc.tables[0].cell(1,0).text == 'Máy chủ'
wb = load_workbook(root/'result.xlsx', data_only=False)
assert wb['Chi phí']['B4'].value == '=SUM(B2:B3)'
assert wb['Chi phí']['B3'].value == 450000
wb.close()
assert 'OmniVoice' in Presentation(root/'result.pptx').slides[0].shapes[1].text
assert '1.650.000' in PdfReader(root/'result.pdf').pages[0].extract_text()
assert 'Tổng ngân sách' in (root/'result-ocr.md').read_text()
with Image.open(root/'result-page.png') as image:
    image.load()
    assert len(image.convert('L').getcolors(image.width*image.height)) > 2
if '--office-ocr' in sys.argv[2:]:
    print('Verified real Soot Office documents, native OCR and decoded PDF render')
    sys.exit(0)
assert not (root/'denied-voice.wav').exists()
for name in ('result-page.png','contact-sheet.png'):
    with Image.open(root/name) as im:
        im.verify()
data = json.loads(subprocess.check_output(['ffprobe','-v','error','-show_streams','-show_format','-of','json',str(root/'result.mp4')]))
video = next(s for s in data['streams'] if s['codec_type']=='video')
assert (video['width'],video['height']) == (360,640)
assert abs(float(data['format']['duration']) - 28/24) < .1
for name in ('narrated.mp4', 'motion.mp4'):
    if not (root/name).exists():
        continue
    data = json.loads(subprocess.check_output(['ffprobe','-v','error','-show_streams','-show_format','-of','json',str(root/name)]))
    video = next(s for s in data['streams'] if s['codec_type']=='video')
    assert (video['width'],video['height']) == (360,640)
    if name == 'narrated.mp4':
        import wave
        with wave.open(str(root/'narration.wav')) as wav:
            duration = wav.getnframes()/wav.getframerate()
        assert float(data['format']['duration']) >= duration
        assert any(s['codec_type']=='audio' for s in data['streams'])
        spec = json.loads((root/'narrated-timeline.json').read_text())
        assert ' '.join(c['text'] for c in spec['captions']) == 'Xin chào Việt Nam. Soot giúp bạn làm việc với tài liệu và video.'
    else:
        assert abs(float(data['format']['duration']) - 2) < .1
for name in ('narrated-contact.png', 'motion-contact.png'):
    if (root/name).exists():
        with Image.open(root/name) as image:
            image.verify()
print(json.dumps({'verified':True,'files':[{ 'path':p.name,'bytes':p.stat().st_size,
    'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in root.iterdir() if p.is_file() and p.name.startswith('result')]}))
